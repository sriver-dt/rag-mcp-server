"""加载器测试：txt/md 编码回退、pdf/docx/html 解析、错误处理。"""

from __future__ import annotations

from pathlib import Path

import pytest

from rag_mcp_server.common.exceptions import LoaderError
from rag_mcp_server.loaders import SUPPORTED_EXTENSIONS, load_document


def test_supported_extensions() -> None:
    assert set(SUPPORTED_EXTENSIONS) == {".txt", ".md", ".pdf", ".docx", ".html", ".htm"}


def test_load_txt(tmp_path: Path) -> None:
    f = tmp_path / "a.txt"
    f.write_text("你好，世界。", encoding="utf-8")
    doc = load_document(f)
    assert doc.text == "你好，世界。"
    assert doc.metadata["filename"] == "a.txt"
    assert doc.metadata["extension"] == ".txt"
    assert doc.metadata["source"] == str(f.resolve())


def test_load_txt_gb18030_fallback(tmp_path: Path) -> None:
    """UTF-8 解码失败时回退 gb18030。"""
    f = tmp_path / "gbk.txt"
    f.write_bytes("中文编码回退测试。".encode("gb18030"))
    assert "中文编码回退测试" in load_document(f).text


def test_load_md(tmp_path: Path) -> None:
    f = tmp_path / "b.md"
    f.write_text("# 标题\n\n正文", encoding="utf-8")
    assert "# 标题" in load_document(f).text


def test_missing_file_raises(tmp_path: Path) -> None:
    with pytest.raises(LoaderError, match="文件不存在"):
        load_document(tmp_path / "no.txt")


def test_unsupported_extension(tmp_path: Path) -> None:
    f = tmp_path / "x.xyz"
    f.write_text("data", encoding="utf-8")
    with pytest.raises(LoaderError, match="不支持的文件类型"):
        load_document(f)


def test_load_pdf(tmp_path: Path) -> None:
    from pypdf import PdfWriter

    f = tmp_path / "doc.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    with f.open("wb") as fh:
        writer.write(fh)

    doc = load_document(f)
    assert doc.metadata["page_count"] == 1
    assert doc.metadata["title"] == "doc"  # 无元数据标题时回退文件名
    assert doc.text == ""  # 空白页无文本


def test_load_docx_preserves_order(tmp_path: Path) -> None:
    """回归：段落与表格按文档原始顺序（段-表-段）输出。"""
    from docx import Document

    f = tmp_path / "t.docx"
    d = Document()
    d.add_paragraph("第一段")
    table = d.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "A1"
    table.cell(0, 1).text = "B1"
    d.add_paragraph("第三段")
    d.save(str(f))

    doc = load_document(f)
    parts = doc.text.split("\n\n")
    assert parts[0] == "第一段"
    assert "A1\tB1" in parts[1]  # 表格出现在原位
    assert parts[2] == "第三段"
    assert doc.metadata["paragraph_count"] == 2
    assert doc.metadata["table_count"] == 1


def test_load_docx_garbage_raises(tmp_path: Path) -> None:
    f = tmp_path / "bad.docx"
    f.write_bytes(b"not a docx")
    with pytest.raises(LoaderError, match="Word 文档打开失败"):
        load_document(f)


def test_load_html_strips_nav_keeps_header(tmp_path: Path) -> None:
    """nav/script 剥离；header 有意保留（含 h1 主标题）。"""
    f = tmp_path / "page.html"
    f.write_text(
        "<html><head><title>页面标题</title></head><body>"
        "<nav>导航链接</nav>"
        "<header><h1>文章主标题</h1></header>"
        "<script>var x = 1;</script>"
        "<p>正文内容</p>"
        "</body></html>",
        encoding="utf-8",
    )
    doc = load_document(f)
    assert doc.metadata["title"] == "页面标题"
    assert "文章主标题" in doc.text
    assert "正文内容" in doc.text
    assert "导航链接" not in doc.text
    assert "var x" not in doc.text
